import sys
from pathlib import Path
import pytest
import openpyxl

SRC_DIR = Path(__file__).resolve().parent.parent / "src"
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

from fastapi.testclient import TestClient
import database
from salary_engine import calculate_weekly_salary
from excel_generator import generate_weekly_closing, WEEKLY_CLOSING_TEMPLATE
from api.main import app

client = TestClient(app)


def setup_comprehensive_test_week():
    conn = database.connect_database()
    cursor = conn.cursor()

    # Create/Get week
    cursor.execute("SELECT id FROM weeks WHERE week_start = '2026-08-01' LIMIT 1")
    row = cursor.fetchone()
    if row:
        week_id = row[0]
    else:
        cursor.execute("INSERT INTO weeks (week_start, week_end) VALUES ('2026-08-01', '2026-08-07')")
        week_id = cursor.lastrowid

    # 1. Department & Contractor
    cursor.execute("SELECT id FROM departments WHERE name = 'POWERTABLE' LIMIT 1")
    dept_row = cursor.fetchone()
    dept_id = dept_row[0] if dept_row else 1

    cursor.execute("SELECT id FROM contractors WHERE name = 'SRI CONTRACTOR' LIMIT 1")
    c_row = cursor.fetchone()
    if not c_row:
        cursor.execute("INSERT INTO contractors (name, commission_amount, status) VALUES ('SRI CONTRACTOR', 40.0, 'active')")
        contractor_id = cursor.lastrowid
    else:
        contractor_id = c_row[0]

    # 2. Employees
    # Piece worker
    cursor.execute("SELECT id FROM employees WHERE employee_code = 'E101' OR name = 'RAMESH' LIMIT 1")
    e_p = cursor.fetchone()
    if not e_p:
        cursor.execute(
            "INSERT INTO employees (employee_code, name, department_id, pay_type) VALUES ('E101', 'RAMESH', ?, 'PIECE')",
            (dept_id,)
        )
        emp_p_id = cursor.lastrowid
    else:
        emp_p_id = e_p[0]

    # Shift worker (Company)
    cursor.execute("SELECT id FROM employees WHERE employee_code = 'E102' OR name = 'SURESH' LIMIT 1")
    e_s1 = cursor.fetchone()
    if not e_s1:
        cursor.execute(
            "INSERT INTO employees (employee_code, name, department_id, pay_type, shift_rate) VALUES ('E102', 'SURESH', ?, 'SHIFT', 500.0)",
            (dept_id,)
        )
        emp_s1_id = cursor.lastrowid
    else:
        emp_s1_id = e_s1[0]

    # Shift worker (Contractor)
    cursor.execute("SELECT id FROM employees WHERE employee_code = 'E103' OR name = 'KUMAR' LIMIT 1")
    e_s2 = cursor.fetchone()
    if not e_s2:
        cursor.execute(
            "INSERT INTO employees (employee_code, name, department_id, pay_type, shift_rate, contractor_id) VALUES ('E103', 'KUMAR', ?, 'SHIFT', 450.0, ?)",
            (dept_id, contractor_id)
        )
        emp_s2_id = cursor.lastrowid
    else:
        emp_s2_id = e_s2[0]


    # GOVINDASAMY E016 (SHIFT)
    cursor.execute("UPDATE employees SET pay_type = 'SHIFT' WHERE employee_code = 'E016'")
    cursor.execute("SELECT id FROM employees WHERE employee_code = 'E016' LIMIT 1")
    gov_row = cursor.fetchone()
    gov_id = gov_row[0] if gov_row else None

    # 3. Master records: Styles & Operations
    cursor.execute("SELECT id FROM styles WHERE style_no = 'STYLE-A' LIMIT 1")
    st1 = cursor.fetchone()
    if not st1:
        cursor.execute("INSERT INTO styles (style_no, style_name) VALUES ('STYLE-A', 'Style A')")
        style1_id = cursor.lastrowid
    else:
        style1_id = st1[0]

    cursor.execute("SELECT id FROM styles WHERE style_no = 'STYLE-B' LIMIT 1")
    st2 = cursor.fetchone()
    if not st2:
        cursor.execute("INSERT INTO styles (style_no, style_name) VALUES ('STYLE-B', 'Style B')")
        style2_id = cursor.lastrowid
    else:
        style2_id = st2[0]


    cursor.execute("SELECT id FROM operations LIMIT 1")
    op1_id = cursor.fetchone()[0]

    # 4. Clean previous records for week
    cursor.execute("DELETE FROM production_sizes WHERE production_record_id IN (SELECT id FROM production_records WHERE week_id = ?)", (week_id,))
    cursor.execute("DELETE FROM production_records WHERE week_id = ?", (week_id,))
    cursor.execute("DELETE FROM shift_records WHERE week_id = ?", (week_id,))
    cursor.execute("DELETE FROM bonuses WHERE week_id = ?", (week_id,))
    cursor.execute("DELETE FROM deductions WHERE week_id = ?", (week_id,))
    cursor.execute("DELETE FROM expenses WHERE week_id = ?", (week_id,))
    cursor.execute("DELETE FROM outsource_payments WHERE week_id = ?", (week_id,))
    cursor.execute("DELETE FROM security_payments WHERE week_id = ?", (week_id,))

    # Insert Production (Multiple styles & items)
    cursor.execute(
        """
        INSERT INTO production_records (week_id, work_date, employee_id, style_id, operation_id, type, total_qty, rate, total_amount)
        VALUES (?, '2026-08-01', ?, ?, ?, 'SHIRT', 50.0, 10.0, 500.0)
        """,
        (week_id, emp_p_id, style1_id, op1_id)
    )
    cursor.execute(
        """
        INSERT INTO production_records (week_id, work_date, employee_id, style_id, operation_id, type, total_qty, rate, total_amount)
        VALUES (?, '2026-08-02', ?, ?, ?, 'PANT', 20.0, 15.0, 300.0)
        """,
        (week_id, emp_p_id, style2_id, op1_id)
    )

    # Insert Shift records (Company & Contractor & Govindasamy)
    cursor.execute(
        """
        INSERT INTO shift_records (week_id, work_date, employee_id, operation_id, item, shifts, shift_rate, daily_salary)
        VALUES (?, '2026-08-01', ?, ?, 'SEWING', 2.0, 500.0, 1000.0)
        """,
        (week_id, emp_s1_id, op1_id)
    )
    cursor.execute(
        """
        INSERT INTO shift_records (week_id, work_date, employee_id, operation_id, item, shifts, shift_rate, daily_salary)
        VALUES (?, '2026-08-01', ?, ?, 'CUTTING', 3.0, 450.0, 1350.0)
        """,
        (week_id, emp_s2_id, op1_id)
    )
    if gov_id:
        cursor.execute(
            """
            INSERT INTO shift_records (week_id, work_date, employee_id, operation_id, item, shifts, shift_rate, daily_salary)
            VALUES (?, '2026-08-01', ?, ?, 'CHECKING', 1.0, 400.0, 400.0)
            """,
            (week_id, gov_id, op1_id)
        )

    # Bonuses & Deductions
    cursor.execute("INSERT INTO bonuses (week_id, employee_id, bonus_type, amount) VALUES (?, ?, 'ATTENDANCE', 150.0)", (week_id, emp_p_id))
    cursor.execute("INSERT INTO deductions (week_id, employee_id, deduction_type, amount) VALUES (?, ?, 'ADVANCE', 50.0)", (week_id, emp_p_id))

    # Expenses
    cursor.execute("INSERT INTO expenses (week_id, category, amount, notes) VALUES (?, 'TEA & SNACKS', 300.0, 'Weekly tea')", (week_id,))
    cursor.execute("INSERT INTO expenses (week_id, category, amount, notes) VALUES (?, 'THREAD', 450.0, 'Sewing thread')", (week_id,))

    # Outsource
    cursor.execute(
        """
        INSERT INTO outsource_payments (week_id, centre_name, style, item, qty, rate, amount)
        VALUES (?, 'EMBROIDERY HUB', 'STYLE-A', 'LOGO EMB', 100.0, 5.0, 500.0)
        """,
        (week_id,)
    )

    # Security
    cursor.execute(
        "INSERT INTO security_payments (week_id, employee_id, days, amount) VALUES (?, ?, 7, 700.0)",
        (week_id, emp_s1_id)
    )

    conn.commit()
    conn.close()
    return week_id


def test_weekly_closing_generation_and_openpyxl_validation():
    week_id = setup_comprehensive_test_week()

    # 1. Salary Engine Calculation
    salary_res = calculate_weekly_salary(week_id)
    assert salary_res is not None

    # 2. Excel Generator Output
    out_file = generate_weekly_closing(week_id, WEEKLY_CLOSING_TEMPLATE)
    assert out_file.exists()
    assert out_file.name == f"Weekly_Closing_Week_{week_id}.xlsx"

    # 3. Programmatically inspect with Openpyxl
    wb = openpyxl.load_workbook(out_file)

    # Sheets exist and names preserved
    expected_sheets = ["TALRS-HELPERS", "PC RATE-OTHERS", "SALARY", "BANK-TRANSFER", "PT-SHIFT"]
    assert set(expected_sheets).issubset(set(wb.sheetnames))

    # Verify no corruption when reading all sheets
    total_cells_read = 0
    for name in expected_sheets:
        ws = wb[name]
        for row in ws.iter_rows(values_only=True):
            total_cells_read += len([c for c in row if c is not None])
    assert total_cells_read > 0

    # 4. Verify API Endpoint POST /weeks/{week_id}/generate-closing
    res = client.post(f"/weeks/{week_id}/generate-closing")
    assert res.status_code == 200
    api_data = res.json()
    assert api_data["week_id"] == week_id
    assert api_data["total_salary"] == salary_res["total_salary"]
    assert api_data["total_contractor_commission"] == salary_res["total_contractor_commission"]
    assert api_data["total_factory_expense"] == salary_res["total_factory_expense"]


def test_invalid_week_id_generate_closing_api():
    res = client.post("/weeks/999999/generate-closing")
    assert res.status_code == 404
    assert "not found" in res.json()["detail"].lower()


if __name__ == "__main__":
    pytest.main(["-v", __file__])
